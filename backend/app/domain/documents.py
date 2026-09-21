"""Deterministic, traceable extraction for user-supplied JD and resume files."""

from __future__ import annotations

import hashlib
import re
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Literal

from docx import Document
from pydantic import BaseModel, Field
from pypdf import PdfReader

MAX_DOCUMENT_BYTES = 5 * 1024 * 1024
DocumentKind = Literal["pdf", "docx", "text", "image"]

IMAGE_MEDIA_TYPES = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".webp": "image/webp",
}


class DocumentExtractionError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        self.code = code
        super().__init__(message)


class ExtractedPage(BaseModel):
    page: int | None
    start_character: int = Field(ge=0)
    end_character: int = Field(ge=0)
    text: str


class ExtractedDocument(BaseModel):
    schema_version: Literal["document-text-v1"] = "document-text-v1"
    filename: str
    document_kind: DocumentKind
    media_type: str
    sha256: str = Field(pattern=r"^[a-f0-9]{64}$")
    extractor_name: str
    extractor_version: Literal["1.0.0"] = "1.0.0"
    text: str = Field(min_length=1, max_length=200_000)
    pages: list[ExtractedPage] = Field(min_length=1, max_length=100)


def _safe_filename(filename: str) -> str:
    name = Path(filename).name.strip()
    if not name or name in {".", ".."}:
        raise DocumentExtractionError("invalid_filename", "A valid filename is required.")
    return name[:255]


def _detect_kind(filename: str, content: bytes) -> DocumentKind:
    suffix = Path(filename).suffix.casefold()
    if content.startswith(b"%PDF-"):
        if suffix != ".pdf":
            raise DocumentExtractionError(
                "type_mismatch", "The file content is PDF but its extension is not .pdf."
            )
        return "pdf"
    if content.startswith(b"PK"):
        if suffix != ".docx":
            raise DocumentExtractionError(
                "type_mismatch", "The file content is ZIP-based but its extension is not .docx."
            )
        try:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                if "word/document.xml" not in archive.namelist():
                    raise DocumentExtractionError(
                        "invalid_docx", "The uploaded archive is not a valid DOCX document."
                    )
        except zipfile.BadZipFile as exc:
            raise DocumentExtractionError("invalid_docx", "The DOCX archive is corrupt.") from exc
        return "docx"
    detected_image: str | None = None
    if content.startswith(b"\xff\xd8\xff"):
        detected_image = "image/jpeg"
    elif content.startswith(b"\x89PNG\r\n\x1a\n"):
        detected_image = "image/png"
    elif content.startswith(b"RIFF") and content[8:12] == b"WEBP":
        detected_image = "image/webp"
    if detected_image:
        if IMAGE_MEDIA_TYPES.get(suffix) != detected_image:
            raise DocumentExtractionError(
                "type_mismatch",
                "The image content does not match its filename extension.",
            )
        return "image"
    if suffix in IMAGE_MEDIA_TYPES:
        raise DocumentExtractionError("invalid_image", "The uploaded image is invalid or corrupt.")
    if suffix in {".txt", ".md"}:
        return "text"
    raise DocumentExtractionError(
        "unsupported_document",
        "Supported document formats are PDF, DOCX, TXT, MD, JPG, JPEG, PNG, and WEBP.",
    )


def media_type_for_document(filename: str, content: bytes) -> str:
    kind = _detect_kind(filename, content)
    if kind == "pdf":
        return "application/pdf"
    if kind == "docx":
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if kind == "text":
        return "text/plain"
    return IMAGE_MEDIA_TYPES[Path(filename).suffix.casefold()]


def _combine_pages(raw_pages: list[tuple[int | None, str]]) -> tuple[str, list[ExtractedPage]]:
    output: list[str] = []
    pages: list[ExtractedPage] = []
    cursor = 0
    for page_number, raw_text in raw_pages:
        page_text = raw_text.replace("\x00", "").strip()
        if not page_text:
            continue
        if output:
            output.append("\n\n")
            cursor += 2
        start = cursor
        output.append(page_text)
        cursor += len(page_text)
        pages.append(
            ExtractedPage(
                page=page_number,
                start_character=start,
                end_character=cursor,
                text=page_text,
            )
        )
    text = "".join(output)
    if not text.strip():
        raise DocumentExtractionError(
            "no_extractable_text",
            "The document contains no extractable text.",
        )
    if len(text) > 200_000:
        raise DocumentExtractionError(
            "document_text_too_large", "Extracted document text exceeds 200,000 characters."
        )
    return text, pages


def _extract_pdf(content: bytes) -> tuple[str, list[ExtractedPage]]:
    try:
        reader = PdfReader(BytesIO(content), strict=False)
        if reader.is_encrypted and reader.decrypt("") == 0:
            raise DocumentExtractionError(
                "encrypted_pdf", "Password-protected PDFs are not supported."
            )
        if len(reader.pages) > 100:
            raise DocumentExtractionError("too_many_pages", "Documents are limited to 100 pages.")
        raw_pages = [
            (index, page.extract_text() or "")
            for index, page in enumerate(reader.pages, 1)
        ]
    except DocumentExtractionError:
        raise
    except Exception as exc:
        raise DocumentExtractionError("invalid_pdf", "The PDF could not be parsed.") from exc
    return _combine_pages(raw_pages)


def _extract_docx(content: bytes) -> tuple[str, list[ExtractedPage]]:
    try:
        document = Document(BytesIO(content))
        blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                value = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if value:
                    blocks.append(value)
    except Exception as exc:
        raise DocumentExtractionError("invalid_docx", "The DOCX could not be parsed.") from exc
    return _combine_pages([(None, "\n".join(blocks))])


def _extract_text(content: bytes) -> tuple[str, list[ExtractedPage]]:
    try:
        decoded = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentExtractionError(
            "invalid_text_encoding", "Text documents must use UTF-8 encoding."
        ) from exc
    return _combine_pages([(None, decoded)])


def extract_document(
    filename: str,
    content: bytes,
    *,
    ocr_text: str | None = None,
) -> ExtractedDocument:
    safe_name = _safe_filename(filename)
    if not content:
        raise DocumentExtractionError("empty_document", "The uploaded document is empty.")
    if len(content) > MAX_DOCUMENT_BYTES:
        raise DocumentExtractionError(
            "document_too_large", "Documents are limited to 5 MiB."
        )
    kind = _detect_kind(safe_name, content)
    if kind == "pdf":
        try:
            text, pages = _extract_pdf(content)
            extractor = "pypdf"
        except DocumentExtractionError as exc:
            if exc.code != "no_extractable_text" or ocr_text is None:
                raise
            text, pages = _combine_pages([(None, ocr_text)])
            extractor = "gemini-vision-ocr"
        media_type = "application/pdf"
    elif kind == "docx":
        text, pages = _extract_docx(content)
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        extractor = "python-docx"
    elif kind == "text":
        text, pages = _extract_text(content)
        media_type = "text/plain"
        extractor = "utf8-text"
    else:
        if ocr_text is None:
            raise DocumentExtractionError(
                "ocr_required",
                "Image job descriptions require OCR by the configured backend service.",
            )
        text, pages = _combine_pages([(1, ocr_text)])
        media_type = media_type_for_document(safe_name, content)
        extractor = "gemini-vision-ocr"
    return ExtractedDocument(
        filename=safe_name,
        document_kind=kind,
        media_type=media_type,
        sha256=hashlib.sha256(content).hexdigest(),
        extractor_name=extractor,
        text=text,
        pages=pages,
    )


def page_for_offset(document: ExtractedDocument, offset: int) -> int | None:
    for page in document.pages:
        if page.start_character <= offset < page.end_character:
            return page.page
    return None


def iter_nonempty_lines(text: str) -> list[tuple[int, int, str]]:
    lines: list[tuple[int, int, str]] = []
    cursor = 0
    for raw_line in text.splitlines(keepends=True):
        line = raw_line.rstrip("\r\n")
        leading = len(line) - len(line.lstrip())
        trailing = len(line.rstrip())
        if trailing > leading:
            lines.append((cursor + leading, cursor + trailing, line[leading:trailing]))
        cursor += len(raw_line)
    if text and not text.endswith(("\n", "\r")) and not lines:
        match = re.search(r"\S.*\S|\S", text)
        if match:
            lines.append((match.start(), match.end(), match.group(0)))
    return lines
