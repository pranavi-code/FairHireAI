from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from backend.app.api.dependencies import authenticated_repository
from backend.app.api.v1 import documents as documents_api
from backend.app.config import Settings, get_settings
from backend.app.domain.documents import DocumentExtractionError, extract_document
from backend.app.domain.resume import extract_resume_evidence
from backend.app.main import app


class ConsentingRepository:
    def __init__(self, *, denied: set[str] | None = None) -> None:
        self.denied = denied or set()

    def consent_granted(self, consent_type: str, _attempt_id=None) -> bool:  # type: ignore[no-untyped-def]
        return consent_type not in self.denied


def _consenting_repository() -> ConsentingRepository:
    return ConsentingRepository()


def test_text_resume_extraction_retains_exact_spans() -> None:
    content = (
        b"Roopika Example\n"
        b"Technical Skills\n"
        b"Python, FastAPI, PostgreSQL\n"
        b"Projects\n"
        b"- Placement readiness API with audit logs\n"
        b"Certifications\n"
        b"Cloud fundamentals certificate\n"
    )
    document = extract_document("resume.txt", content)
    evidence = extract_resume_evidence(document)

    assert document.sha256 == evidence.document_sha256
    assert [claim.claim_type for claim in evidence.claims] == [
        "skill",
        "skill",
        "skill",
        "project",
        "certification",
    ]
    for claim in evidence.claims:
        source = claim.source
        assert document.text[source.start_character : source.end_character] == source.source_text
    assert evidence.claims[0].normalized_text == "Python"


def test_docx_extraction_reads_paragraphs_and_tables() -> None:
    source = Document()
    source.add_paragraph("Junior Backend Developer")
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "SQL"
    output = BytesIO()
    source.save(output)

    document = extract_document("role.docx", output.getvalue())

    assert document.document_kind == "docx"
    assert "Junior Backend Developer" in document.text
    assert "Python | SQL" in document.text
    assert document.pages[0].page is None


def test_scanned_pdf_accepts_ocr_text() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=600, height=800)
    output = BytesIO()
    writer.write(output)
    ocr_text = "Junior Data Analyst role requiring SQL, Python, dashboards, and statistics."

    document = extract_document("scanned-job.pdf", output.getvalue(), ocr_text=ocr_text)

    assert document.document_kind == "pdf"
    assert document.extractor_name == "gemini-vision-ocr"
    assert document.text == ocr_text


@pytest.mark.parametrize(
    ("filename", "content", "media_type"),
    [
        ("resume.jpg", b"\xff\xd8\xffimage", "image/jpeg"),
        ("resume.jpeg", b"\xff\xd8\xffimage", "image/jpeg"),
        ("resume.png", b"\x89PNG\r\n\x1a\nimage", "image/png"),
    ],
)
def test_requested_image_document_formats_accept_ocr_text(
    filename: str,
    content: bytes,
    media_type: str,
) -> None:
    document = extract_document(filename, content, ocr_text="Skills\nPython and SQL")

    assert document.document_kind == "image"
    assert document.media_type == media_type
    assert document.extractor_name == "gemini-vision-ocr"


def test_extension_and_content_mismatch_is_rejected() -> None:
    try:
        extract_document("fake.txt", b"%PDF-1.7\n")
    except DocumentExtractionError as exc:
        assert exc.code == "type_mismatch"
    else:
        raise AssertionError("Mismatched PDF content should be rejected")


def test_jd_upload_extracts_and_maps_supported_role() -> None:
    app.dependency_overrides[authenticated_repository] = _consenting_repository
    client = TestClient(app)
    jd = (
        "We are hiring a Junior Backend Developer to build REST APIs using "
        "Python, FastAPI, PostgreSQL, Git, testing, and debugging."
    )

    try:
        response = client.post(
            "/api/v1/documents/job-description",
            files={"file": ("job.txt", jd.encode(), "text/plain")},
        )
    finally:
        app.dependency_overrides.pop(authenticated_repository, None)

    assert response.status_code == 200
    payload = response.json()
    assert payload["document"]["sha256"]
    assert payload["role_mapping"]["status"] == "detected"
    assert payload["role_mapping"]["mapping_version"] == "optional-jd-mapper-v2"
    assert payload["role_mapping"]["role_id"] == "junior_backend_developer"


def test_image_jd_uses_configured_backend_ocr(monkeypatch) -> None:  # type: ignore[no-untyped-def]
    image_bytes = b"\xff\xd8\xff" + (b"image-data" * 20)
    ocr_text = (
        "We are hiring a Junior Backend Developer to build REST APIs with Python, "
        "FastAPI, PostgreSQL, testing, Git, and debugging."
    )

    class FakeGeminiClient:
        def __init__(self, **_: object) -> None:
            pass

        def __enter__(self) -> "FakeGeminiClient":
            return self

        def __exit__(self, *_: object) -> None:
            pass

        def extract_text_from_media(self, *, content: bytes, media_type: str) -> str:
            assert content == image_bytes
            assert media_type == "image/jpeg"
            return ocr_text

    monkeypatch.setattr(documents_api, "GeminiClient", FakeGeminiClient)
    app.dependency_overrides[authenticated_repository] = _consenting_repository
    app.dependency_overrides[get_settings] = lambda: Settings(gemini_api_key="test-key")
    client = TestClient(app)
    try:
        response = client.post(
            "/api/v1/documents/job-description",
            files={"file": ("job.jpeg", image_bytes, "image/jpeg")},
        )
    finally:
        app.dependency_overrides.pop(authenticated_repository, None)
        app.dependency_overrides.pop(get_settings, None)

    assert response.status_code == 200
    payload = response.json()
    assert payload["document"]["document_kind"] == "image"
    assert payload["document"]["extractor_name"] == "gemini-vision-ocr"
    assert payload["document"]["text"] == ocr_text
    assert payload["role_mapping"]["role_id"] == "junior_backend_developer"


def test_image_jd_reports_when_ocr_is_not_configured() -> None:
    app.dependency_overrides[authenticated_repository] = _consenting_repository
    app.dependency_overrides[get_settings] = lambda: Settings(gemini_api_key=None)
    client = TestClient(app)
    try:
        response = client.post(
            "/api/v1/documents/job-description",
            files={"file": ("job.png", b"\x89PNG\r\n\x1a\nimage", "image/png")},
        )
    finally:
        app.dependency_overrides.pop(authenticated_repository, None)
        app.dependency_overrides.pop(get_settings, None)

    assert response.status_code == 503
    assert response.json()["detail"]["code"] == "ocr_unavailable"


def test_image_jd_requires_external_ai_consent_before_ocr() -> None:
    app.dependency_overrides[authenticated_repository] = lambda: ConsentingRepository(
        denied={"external_ai_processing"}
    )
    app.dependency_overrides[get_settings] = lambda: Settings(gemini_api_key="test-key")
    client = TestClient(app)
    try:
        response = client.post(
            "/api/v1/documents/job-description",
            files={"file": ("job.png", b"\x89PNG\r\n\x1a\nimage", "image/png")},
        )
    finally:
        app.dependency_overrides.pop(authenticated_repository, None)
        app.dependency_overrides.pop(get_settings, None)

    assert response.status_code == 403
    assert response.json()["detail"]["code"] == "external_ai_processing_consent_required"


def test_resume_upload_returns_traceable_evidence() -> None:
    app.dependency_overrides[authenticated_repository] = _consenting_repository
    client = TestClient(app)
    try:
        response = client.post(
            "/api/v1/documents/resume",
            files={
                "file": (
                    "resume.md",
                    b"Skills\nPython, SQL\nProjects\n- Built a REST API",
                    "text/markdown",
                )
            },
        )
    finally:
        app.dependency_overrides.pop(authenticated_repository, None)

    assert response.status_code == 200
    payload = response.json()
    assert len(payload["evidence"]["claims"]) == 3
    assert payload["evidence"]["claims"][0]["source"]["source_text"] == "Python"


def test_image_resume_uses_ocr_only_with_both_required_consents(monkeypatch) -> None:  # type: ignore[no-untyped-def]
    image_bytes = b"\x89PNG\r\n\x1a\n" + (b"image-data" * 20)
    ocr_text = "Skills\nPython, SQL\nProjects\nBuilt a reliable REST API"

    class FakeGeminiClient:
        def __init__(self, **_: object) -> None:
            pass

        def __enter__(self) -> "FakeGeminiClient":
            return self

        def __exit__(self, *_: object) -> None:
            pass

        def extract_text_from_media(self, *, content: bytes, media_type: str) -> str:
            assert content == image_bytes
            assert media_type == "image/png"
            return ocr_text

    monkeypatch.setattr(documents_api, "GeminiClient", FakeGeminiClient)
    app.dependency_overrides[authenticated_repository] = _consenting_repository
    app.dependency_overrides[get_settings] = lambda: Settings(gemini_api_key="test-key")
    client = TestClient(app)
    try:
        response = client.post(
            "/api/v1/documents/resume",
            files={"file": ("resume.png", image_bytes, "image/png")},
        )
    finally:
        app.dependency_overrides.pop(authenticated_repository, None)
        app.dependency_overrides.pop(get_settings, None)

    assert response.status_code == 200
    assert response.json()["document"]["extractor_name"] == "gemini-vision-ocr"
    assert response.json()["evidence"]["claims"]


def test_resume_extraction_requires_resume_processing_consent() -> None:
    app.dependency_overrides[authenticated_repository] = lambda: ConsentingRepository(
        denied={"resume_processing"}
    )
    client = TestClient(app)
    try:
        response = client.post(
            "/api/v1/documents/resume",
            files={"file": ("resume.txt", b"Skills\nPython", "text/plain")},
        )
    finally:
        app.dependency_overrides.pop(authenticated_repository, None)

    assert response.status_code == 403
    assert response.json()["detail"]["code"] == "resume_processing_consent_required"


def test_scanned_or_blank_text_is_rejected() -> None:
    app.dependency_overrides[authenticated_repository] = _consenting_repository
    client = TestClient(app)
    try:
        response = client.post(
            "/api/v1/documents/resume",
            files={"file": ("empty.txt", b" \n\t", "text/plain")},
        )
    finally:
        app.dependency_overrides.pop(authenticated_repository, None)
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "no_extractable_text"


def test_document_upload_requires_authentication() -> None:
    client = TestClient(app)
    response = client.post(
        "/api/v1/documents/resume",
        files={"file": ("resume.txt", b"Skills\nPython", "text/plain")},
    )
    assert response.status_code == 401
